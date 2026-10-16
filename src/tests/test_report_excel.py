"""گزارش اکسل — ساختار، بازتولیدپذیری، پوشاندن داده و حاکمیت خروجی."""

from __future__ import annotations

from datetime import UTC, datetime, timedelta
from io import BytesIO

from django.utils import timezone
from openpyxl import load_workbook
import pytest

from apps.audit.choices import AuditAction
from apps.audit.models import AuditRecord
from apps.reporting.archive import XLSX_CONTENT_TYPE
from apps.reporting.excel import build_workbook
from apps.reporting.ranges import MAX_SPAN_DAYS, ReportRange, parse_range
from apps.reporting.services import build_report, build_report_for_range
from tests.factories import future_period

pytestmark = [pytest.mark.django_db, pytest.mark.integration]

SHEET_NAMES = [
    "راهنما",
    "رزروها",
    "تقویم سانس‌ها",
    "اشغال واحدها",
    "آمار روزانه",
    "کنش‌های کارکنان",
    "پیامک‌ها",
]


def _fixed_moment() -> datetime:
    """لحظهٔ ثابت برای آزمون بازتولیدپذیری."""
    return datetime(2026, 10, 20, 9, 30, tzinfo=UTC)


def _range() -> ReportRange:
    """بازهٔ ثابت پوشانندهٔ رزروهای تست."""
    today = timezone.localdate()
    return ReportRange(start_day=today, end_day=today + timedelta(days=14))


def _sheet_values(workbook, name: str) -> list[list[object]]:
    """همهٔ مقادیر یک برگه به‌صورت جدول."""
    return [list(row) for row in workbook[name].iter_rows(values_only=True)]


# ── ساختار سند ───────────────────────────────────────────────────────────────


def test_workbook_has_seven_rtl_print_ready_sheets(units, guest) -> None:
    """هفت برگه، همه راست‌به‌چپ، با عنوان تکرارشو در چاپ و پالایهٔ خودکار."""
    workbook = build_workbook(
        __import__("apps.reporting.dataset", fromlist=["build_dataset"]).build_dataset(
            report_range=_range(), generated_at=_fixed_moment()
        )
    )
    assert workbook.sheetnames == SHEET_NAMES
    for name in SHEET_NAMES:
        sheet = workbook[name]
        assert sheet.sheet_view.rightToLeft is True
        assert sheet.sheet_view.showGridLines is False
        assert sheet.page_setup.orientation == "landscape"
    assert workbook["رزروها"].print_title_rows == "$5:$5"  # تکرار سرستون در چاپ
    assert workbook["رزروها"].freeze_panes == "A6"
    assert workbook["آمار روزانه"].auto_filter.ref


def test_reservation_sheet_carries_business_columns(api, guest, units) -> None:
    """برگهٔ رزروها: کد، وضعیت، تاریخ شمسی، نیم‌روز و تلفن پوشانده."""
    start, end = future_period(days=2, cells=2)
    api.force_authenticate(user=guest)
    created = api.post(
        "/api/v1/reservations",
        {
            "unit_code": "A1",
            "start": start.isoformat(),
            "end": end.isoformat(),
            "guests_count": 2,
        },
        format="json",
    )
    assert created.status_code == 201

    report = build_report_for_range(report_range=_range(), generated_at=_fixed_moment())
    workbook = load_workbook(BytesIO(report.content))
    rows = _sheet_values(workbook, "رزروها")
    header = rows[4]
    assert "کد رزرو" in header
    assert "تاریخ ورود (شمسی)" in header or "ورود (شمسی)" in header

    data_rows = [row for row in rows[5:] if row and row[0]]
    assert len(data_rows) == 1
    row = dict(zip([str(item) for item in header], data_rows[0], strict=False))
    assert row["کد رزرو"] == created.json()["code"]
    assert row["واحد"] == "A1"
    assert row["نیم‌روز"] == 2
    assert row["تلفن"].endswith(guest.phone[-4:])
    assert "***" in row["تلفن"]


def test_slots_sheet_covers_every_unit_and_half_day(api, guest, units) -> None:
    """برگهٔ تقویم سانس‌ها: ۵ واحد × ۲ نیم‌روز × تعداد روز."""
    report_range = ReportRange(start_day=timezone.localdate(), end_day=timezone.localdate())
    report = build_report_for_range(report_range=report_range, generated_at=_fixed_moment())
    assert report.rows["slots"] == report_range.days * 5 * 2


def test_phone_numbers_are_masked_everywhere(api, guest, units) -> None:
    """هیچ شمارهٔ کاملی در هیچ برگه‌ای نیست — گزارش هم دادهٔ حساس است."""
    start, end = future_period(days=1, cells=2)
    api.force_authenticate(user=guest)
    api.post(
        "/api/v1/reservations",
        {"unit_code": "B1", "start": start.isoformat(), "end": end.isoformat(), "guests_count": 1},
        format="json",
    )
    report = build_report_for_range(report_range=_range(), generated_at=_fixed_moment())
    workbook = load_workbook(BytesIO(report.content))
    digits = "".join(character for character in guest.phone if character.isdigit())[-10:]
    for name in SHEET_NAMES:
        for row in workbook[name].iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str):
                    assert digits not in value, f"شمارهٔ کامل در برگهٔ «{name}» پیدا شد"


# ── بازتولیدپذیری ────────────────────────────────────────────────────────────


def test_two_identical_reports_are_byte_identical(guest, units) -> None:
    """دو خروجی با داده و زمان یکسان، بایت‌به‌بایت یکسان‌اند (سند قابل ارجاع)."""
    report_range = _range()
    first = build_report_for_range(report_range=report_range, generated_at=_fixed_moment())
    second = build_report_for_range(report_range=report_range, generated_at=_fixed_moment())
    assert first.content == second.content
    assert first.sha256 == second.sha256
    assert first.filename == second.filename == f"mahdieh-report-{report_range.slug}.xlsx"


def test_report_hash_changes_when_data_changes(api, guest, units) -> None:
    """تغییر داده ⇒ تغییر هش (اثر انگشت واقعاً محتوا را می‌سنجد)."""
    before = build_report_for_range(report_range=_range(), generated_at=_fixed_moment())
    start, end = future_period(days=3, cells=2)
    api.force_authenticate(user=guest)
    api.post(
        "/api/v1/reservations",
        {"unit_code": "A2", "start": start.isoformat(), "end": end.isoformat(), "guests_count": 2},
        format="json",
    )
    after = build_report_for_range(report_range=_range(), generated_at=_fixed_moment())
    assert before.sha256 != after.sha256
    assert after.rows["reservations"] == 1


def test_report_is_registered_in_the_audit_chain(operator, units) -> None:
    """هر خروجی گزارش، یک رکورد حسابرسی با هش و بازه دارد."""
    report = build_report_for_range(
        report_range=_range(), generated_at=_fixed_moment(), actor=operator
    )
    record = AuditRecord.all_objects.filter(action=AuditAction.EXPORT_REQUESTED).first()
    assert record is not None
    assert record.metadata["file_sha256"] == report.sha256
    assert record.metadata["rows"]["slots"] == report.rows["slots"]


# ── بازهٔ گزارش ──────────────────────────────────────────────────────────────


def test_parse_range_defaults_and_ordering() -> None:
    """پیش‌فرض یک ماه اخیر، و بازهٔ معکوس مرتب می‌شود."""
    today = timezone.localdate()
    default = parse_range(from_raw=None, to_raw=None, today=today)
    assert default.end_day == today
    assert default.days == 30

    reversed_range = parse_range(from_raw="2026-10-20", to_raw="2026-10-01", today=today)
    assert reversed_range.start_day.isoformat() == "2026-10-01"
    assert reversed_range.end_day.isoformat() == "2026-10-20"


def test_parse_range_rejects_bad_input() -> None:
    """ورودی نامعتبر و بازهٔ بسیار بلند رد می‌شوند."""
    from apps.core.errors import ValidationError as DomainValidationError

    with pytest.raises(DomainValidationError) as bad_date:
        parse_range(from_raw="1405/07/01")
    assert bad_date.value.code == "invalid_date"

    assert MAX_SPAN_DAYS == 366
    with pytest.raises(DomainValidationError) as too_long:
        parse_range(from_raw="2020-01-01", to_raw="2026-01-01")
    assert too_long.value.code == "range_too_long"


# ── مسیر HTTP ────────────────────────────────────────────────────────────────


def test_report_endpoint_is_staff_only(api, units) -> None:
    """ناشناس ۴۰۱، زائر ۴۰۳ (خروجی گزارش دادهٔ عملیاتی حساس است)."""
    from tests.factories import make_user

    assert api.get("/api/v1/staff/reports/reservations.xlsx").status_code == 401
    api.force_authenticate(user=make_user(phone="09120000091"))
    assert api.get("/api/v1/staff/reports/reservations.xlsx").status_code == 403


def test_report_endpoint_returns_named_masked_workbook(api, operator, guest, units) -> None:
    """پاسخ: فایل اکسل با نام، هدر هش، و کارپوشهٔ سالم و کامل."""
    start, end = future_period(days=2, cells=2)
    api.force_authenticate(user=guest)
    api.post(
        "/api/v1/reservations",
        {"unit_code": "B3", "start": start.isoformat(), "end": end.isoformat(), "guests_count": 2},
        format="json",
    )

    api.force_authenticate(user=operator)
    today = timezone.localdate()
    response = api.get(
        "/api/v1/staff/reports/reservations.xlsx",
        {"from": today.isoformat(), "to": (today + timedelta(days=14)).isoformat()},
    )
    assert response.status_code == 200
    assert response["Content-Type"] == XLSX_CONTENT_TYPE
    assert response["Content-Disposition"].startswith('attachment; filename="mahdieh-report-')
    assert len(response["X-Report-SHA256"]) == 64
    assert int(response["X-Report-Rows"]) > 0

    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == SHEET_NAMES
    assert workbook["راهنما"]["B1"].value


def test_report_summary_endpoint_gives_numbers_first(api, operator, guest, units) -> None:
    """خلاصهٔ پیش از دانلود: شمار رزروها، نرخ اشغال و اثر انگشت محتوا."""
    start, end = future_period(days=4, cells=2)
    api.force_authenticate(user=guest)
    api.post(
        "/api/v1/reservations",
        {"unit_code": "A1", "start": start.isoformat(), "end": end.isoformat(), "guests_count": 3},
        format="json",
    )
    api.force_authenticate(user=operator)
    today = timezone.localdate()
    response = api.get(
        "/api/v1/staff/reports/summary",
        {"from": today.isoformat(), "to": (today + timedelta(days=14)).isoformat()},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["totals"]["reservations"] == 1
    assert payload["rows"]["slots"] > 0
    assert len(payload["payload_digest"]) == 64


def test_report_rejects_unknown_unit_and_long_range(api, operator, units) -> None:
    """کد واحد ناشناخته و بازهٔ بیش از یک‌ساله ⇒ ۴۲۲ با کد روشن."""
    api.force_authenticate(user=operator)
    unknown = api.get("/api/v1/staff/reports/reservations.xlsx", {"units": "ZZ"})
    assert unknown.status_code == 422
    assert unknown.json()["code"] == "unknown_unit"

    too_long = api.get(
        "/api/v1/staff/reports/reservations.xlsx",
        {"from": "2020-01-01", "to": "2026-01-01"},
    )
    assert too_long.status_code == 422
    assert too_long.json()["code"] == "range_too_long"


def test_build_report_accepts_raw_query_strings(operator, units) -> None:
    """سرویس، بازهٔ خام رشته‌ای را هم می‌پذیرد (مسیر دستورهای مدیریتی)."""
    report = build_report(from_raw=None, to_raw=None, actor=operator)
    assert report.content[:2] == b"PK"  # آرشیو ZIP استاندارد اکسل
    assert report.filename.endswith(".xlsx")
