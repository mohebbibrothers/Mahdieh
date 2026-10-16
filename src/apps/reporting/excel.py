"""سازندهٔ کارپوشهٔ اکسل — هفت برگه، راست‌به‌چپ، آمادهٔ چاپ و بازتولیدپذبر.

اصول:
    • **خوانایی پیش از تزئین:** هر برگه یک بلوک عنوان، یک سرستون یخ‌زده،
      پالایهٔ خودکار و ستون‌های عرض‌متناسب دارد.
    • **بی‌رنگی به‌جای شلوغی:** رنگ فقط معنا می‌بخشد (وضعیت، نرخ اشغال).
    • **بازتولیدپذیری:** هیچ تاریخی جز ``generated_at`` در فایل نیست و متادادهٔ
      سند ثابت است؛ پس دو اجرا با دادهٔ یکسان، فایل یکسان می‌دهد.
"""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from apps.reporting import theme
from apps.reporting.dataset import ReportData
from apps.reporting.ranges import ReportRange
from apps.reporting.theme import Column

TITLE_ROW = 1
SUBTITLE_ROW = 2
META_ROW = 3
HEADER_ROW = 5
FIRST_DATA_ROW = HEADER_ROW + 1

_THIN = Side(style="thin", color=theme.BORDER)


def build_workbook(data: ReportData) -> Workbook:
    """ساخت کارپوشهٔ کامل گزارش از دادهٔ آماده."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    _stamp(workbook, data.generated_at)

    _guide_sheet(workbook, data)
    _table_sheet(
        workbook,
        name="رزروها",
        title="فهرست رزروها",
        subtitle="هر ردیف یک رزرو؛ تاریخ‌ها شمسی و تلفن پوشانده است.",
        columns=_RESERVATION_COLUMNS,
        rows=data.reservations,
    )
    _table_sheet(
        workbook,
        name="تقویم سانس‌ها",
        title="تقویم سانس‌ها",
        subtitle="وضعیت هر نیم‌روز هر واحد: آزاد، اشغال، بافر نظافت، غیرقابل‌رزرو.",
        columns=_SLOT_COLUMNS,
        rows=data.slots,
    )
    _table_sheet(
        workbook,
        name="اشغال واحدها",
        title="کارنامهٔ اشغال واحدها",
        subtitle="شمار نیم‌روزها و نرخ اشغال به‌ازای هر واحد در بازهٔ گزارش.",
        columns=_UNIT_COLUMNS,
        rows=data.units,
        occupancy_keys=("occupancy",),
    )
    _table_sheet(
        workbook,
        name="آمار روزانه",
        title="آمار روزانه",
        subtitle="نبض روزبه‌روز: اشغال، ورودی، خروجی و سطح تقاضا.",
        columns=_DAILY_COLUMNS,
        rows=data.daily,
        occupancy_keys=("occupancy",),
    )
    _table_sheet(
        workbook,
        name="کنش‌های کارکنان",
        title="کنش‌های کارکنان",
        subtitle="پذیرش، تخلیه، لغو و جابه‌جایی — با نام کارمند و دلیل ثبت‌شده.",
        columns=_ACTION_COLUMNS,
        rows=data.staff_actions,
    )
    _table_sheet(
        workbook,
        name="پیامک‌ها",
        title="گزارش پیامک‌ها",
        subtitle="ارسال‌ها با شمارهٔ پوشانده؛ برای ممیزی و پیگیری تحویل.",
        columns=_SMS_COLUMNS,
        rows=data.sms,
    )
    return workbook


def _stamp(workbook: Workbook, generated_at: datetime) -> None:
    """ثبت متادادهٔ ثابت سند — شرط لازم بازتولیدپذیری بایت‌به‌بایت."""
    properties = workbook.properties
    properties.title = theme.DOCUMENT_TITLE
    properties.subject = theme.DOCUMENT_SUBTITLE
    properties.creator = theme.DOCUMENT_FOOTER
    properties.lastModifiedBy = theme.DOCUMENT_FOOTER
    properties.description = "خروجی رسمی سامانهٔ رزرو زائرسرای مهدیه"
    properties.category = "گزارش عملیاتی"
    properties.keywords = "مهدیه, رزرو, زائرسرا, گزارش"
    properties.created = generated_at
    properties.modified = generated_at


# ── برگهٔ راهنما ─────────────────────────────────────────────────────────────


def _guide_sheet(workbook: Workbook, data: ReportData) -> None:
    """برگهٔ نخست: هویت سند، پالایه‌ها، جمع‌بندی و اثر انگشت محتوا."""
    sheet = workflow_sheet(workbook, "راهنما")
    sheet.column_dimensions["A"].width = 3
    sheet.column_dimensions["B"].width = 34
    sheet.column_dimensions["C"].width = 52

    sheet.cell(row=1, column=2, value=theme.DOCUMENT_TITLE).font = Font(
        name=theme.FONT_FAMILY, size=18, bold=True, color=theme.BRAND_DARK
    )
    sheet.cell(row=2, column=2, value=theme.DOCUMENT_SUBTITLE).font = Font(
        name=theme.FONT_FAMILY, size=11, color=theme.MUTED
    )
    sheet.row_dimensions[1].height = 30

    report_range = data.report_range
    totals = data.totals
    rows: list[tuple[str, Any]] = [
        ("بازهٔ گزارش (شمسی)", report_range.label),
        ("بازهٔ گزارش (میلادی)", f"{report_range.start_day} → {report_range.end_day}"),
        ("تعداد روزها", report_range.days),
        ("زمان تولید گزارش", f"{_jalali_stamp(data.generated_at)} (شمسی)"),
        ("", ""),
        ("شمار رزروها", totals.get("reservations", 0)),
        ("رزروهای فعال", totals.get("active", 0)),
        ("رزروهای پایانیافته/لغوشده", totals.get("cancelled", 0)),
        ("مجموع نفرات", totals.get("guests", 0)),
        ("نیم‌روزهای کل", totals.get("total_cells", 0)),
        ("نیم‌روزهای اشغال", totals.get("occupied_cells", 0)),
        ("نرخ اشغال بازه", f"{(totals.get('occupancy') or 0) * 100:.1f}٪"),
        ("", ""),
        ("ردیف‌های «رزروها»", data.row_counts["reservations"]),
        ("ردیف‌های «تقویم سانس‌ها»", data.row_counts["slots"]),
        ("ردیف‌های «کنش‌های کارکنان»", data.row_counts["staff_actions"]),
        ("ردیف‌های «پیامک‌ها»", data.row_counts["sms"]),
        ("", ""),
        ("اثر انگشت محتوا (SHA-256)", data.payload_digest()),
        ("نسخهٔ سند", theme.DOCUMENT_VERSION),
    ]

    row_index = 4
    for label, value in rows:
        if label:
            label_cell = sheet.cell(row=row_index, column=2, value=label)
            label_cell.font = Font(name=theme.FONT_FAMILY, size=11, bold=True, color=theme.TEXT)
            label_cell.fill = PatternFill("solid", start_color=theme.SURFACE)
            label_cell.alignment = Alignment(horizontal="right", vertical="center")
            value_cell = sheet.cell(row=row_index, column=3, value=value)
            value_cell.font = Font(name=theme.FONT_FAMILY, size=11, color=theme.TEXT)
            value_cell.alignment = Alignment(horizontal="right", vertical="center")
            value_cell.border = Border(bottom=_THIN)
        row_index += 1

    by_status = totals.get("by_status", {})
    row_index += 1
    sheet.cell(row=row_index, column=2, value="تفکیک وضعیت رزروها").font = Font(
        name=theme.FONT_FAMILY, size=12, bold=True, color=theme.BRAND
    )
    row_index += 1
    for status_label, count in by_status.items():
        sheet.cell(row=row_index, column=2, value=status_label).font = Font(
            name=theme.FONT_FAMILY, size=11, color=theme.TEXT
        )
        sheet.cell(row=row_index, column=3, value=count).font = Font(
            name=theme.FONT_FAMILY, size=11, bold=True, color=theme.BRAND_DARK
        )
        row_index += 1

    footer = sheet.cell(row=row_index + 2, column=2, value=theme.DOCUMENT_FOOTER)
    footer.font = Font(name=theme.FONT_FAMILY, size=10, italic=True, color=theme.MUTED)


def workflow_sheet(workbook: Workbook, name: str) -> Worksheet:
    """ساخت برگهٔ راست‌به‌چپ با تنظیمات چاپ حرفه‌ای (A4 افقی، عنوان تکرارشو)."""
    sheet = workbook.create_sheet(title=name)
    sheet.sheet_view.rightToLeft = True
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = theme.BRAND
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    return sheet


# ── برگه‌های جدولی ───────────────────────────────────────────────────────────


def _table_sheet(
    workbook: Workbook,
    *,
    name: str,
    title: str,
    subtitle: str,
    columns: Sequence[Column],
    rows: Iterable[dict[str, Any]],
    occupancy_keys: tuple[str, ...] = (),
) -> None:
    """ساخت یک برگهٔ جدولی کامل: عنوان، سرستون، داده، پالایه و قالب‌بندی."""
    sheet = workflow_sheet(workbook, name)
    span = max(len(columns), 1)
    sheet.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=span)
    sheet.cell(row=TITLE_ROW, column=1, value=title).font = Font(
        name=theme.FONT_FAMILY, size=15, bold=True, color=theme.BRAND_DARK
    )
    sheet.merge_cells(start_row=SUBTITLE_ROW, start_column=1, end_row=SUBTITLE_ROW, end_column=span)
    sheet.cell(row=SUBTITLE_ROW, column=1, value=subtitle).font = Font(
        name=theme.FONT_FAMILY, size=10, color=theme.MUTED
    )
    sheet.row_dimensions[TITLE_ROW].height = 26

    _write_header(sheet, columns)
    row_index = _write_rows(sheet, columns, rows)
    _style_columns(sheet, columns, row_index)
    if row_index >= FIRST_DATA_ROW:
        sheet.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(span)}{row_index}"
    sheet.freeze_panes = f"A{FIRST_DATA_ROW}"
    for key in occupancy_keys:
        _occupancy_rule(sheet, columns, key, last_row=row_index)


def _write_header(sheet: Worksheet, columns: Sequence[Column]) -> None:
    """نوشتن سرستون‌های رنگ‌شده و یخ‌زده."""
    fill = PatternFill("solid", start_color=theme.BRAND)
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=index, value=column.header)
        cell.font = Font(name=theme.FONT_FAMILY, size=11, bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=theme.BRAND_DARK))
    sheet.row_dimensions[HEADER_ROW].height = 28


def _write_rows(sheet: Worksheet, columns: Sequence[Column], rows: Iterable[dict[str, Any]]) -> int:
    """نوشتن ردیف‌های داده با راه‌راه ملایم و بازگرداندن آخرین شمارهٔ ردیف."""
    zebra = PatternFill("solid", start_color=theme.ZEBRA)
    row_index = FIRST_DATA_ROW
    for offset, row in enumerate(rows):
        for column_index, column in enumerate(columns, start=1):
            value = row.get(column.key, "")
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.font = Font(name=theme.FONT_FAMILY, size=10, color=theme.TEXT)
            cell.alignment = Alignment(
                horizontal="center" if column.align == "center" else "right", vertical="center"
            )
            cell.border = Border(bottom=_THIN)
            if column.number_format:
                cell.number_format = column.number_format
            if column.key in {"status", "status_label"} and isinstance(value, str):
                cell.font = Font(
                    name=theme.FONT_FAMILY,
                    size=10,
                    bold=column.key == "status_label",
                    color=theme.status_color(row.get("status", ""))
                    if column.key == "status_label"
                    else theme.TEXT,
                )
            if column.key == "occupancy" and isinstance(value, int | float):
                cell.font = Font(
                    name=theme.FONT_FAMILY,
                    size=10,
                    bold=True,
                    color=theme.occupancy_color(float(value)),
                )
            if offset % 2 == 1:
                cell.fill = zebra
        row_index += 1
    return row_index - 1


def _style_columns(sheet: Worksheet, columns: Sequence[Column], last_row: int) -> None:
    """تعیین عرض ستون‌ها و بلندی ثابت ردیف‌ها برای خوانایی چاپ."""
    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = column.width
    for row_index in range(FIRST_DATA_ROW, last_row + 1):
        sheet.row_dimensions[row_index].height = 20


def _occupancy_rule(
    sheet: Worksheet, columns: Sequence[Column], key: str, *, last_row: int
) -> None:
    """قالب‌بندی شرطی نرخ اشغال — نگاه یک‌نگاهی به فشار ظرفیت."""
    if last_row < FIRST_DATA_ROW:
        return
    index = next(
        (position for position, column in enumerate(columns, 1) if column.key == key), None
    )
    if index is None:
        return
    letter = get_column_letter(index)
    sheet.conditional_formatting.add(
        f"{letter}{FIRST_DATA_ROW}:{letter}{last_row}",
        ColorScaleRule(
            start_type="num",
            start_value=0,
            start_color="E8F5E9",
            mid_type="num",
            mid_value=0.6,
            mid_color="FFF3CD",
            end_type="num",
            end_value=1,
            end_color="F8D7DA",
        ),
    )


# ── تعریف ستون‌ها ────────────────────────────────────────────────────────────

_RESERVATION_COLUMNS: tuple[Column, ...] = (
    Column("کد رزرو", "code", 13, align="center"),
    Column("وضعیت", "status_label", 14, align="center"),
    Column("منبع", "source", 10, align="center"),
    Column("طبقه", "floor", 7, align="center"),
    Column("واحد", "unit", 8, align="center"),
    Column("عنوان واحد", "unit_title", 20),
    Column("نوع واحد", "unit_type", 12, align="center"),
    Column("ورود (شمسی)", "start_jalali", 13, align="center"),
    Column("ساعت ورود", "start_time", 10, align="center"),
    Column("خروج (شمسی)", "end_jalali", 13, align="center"),
    Column("ساعت خروج", "end_time", 10, align="center"),
    Column("نیم‌روز", "cells", 9, align="center"),
    Column("مدت اقامت", "duration", 20),
    Column("نفرات", "guests", 8, align="center"),
    Column("نام زائر", "guest_name", 20),
    Column("تلفن", "guest_phone", 15, align="center"),
    Column("استان", "province", 14),
    Column("شهر", "city", 14),
    Column("ثبت‌کننده", "created_by", 18),
    Column("یادداشت زائر", "guest_note", 26),
    Column("یادداشت داخلی", "staff_note", 26),
    Column("زمان ثبت", "created_at", 17, align="center"),
)

_SLOT_COLUMNS: tuple[Column, ...] = (
    Column("تاریخ", "day", 12, align="center"),
    Column("تاریخ شمسی", "date_jalali", 13, align="center"),
    Column("روز هفته", "weekday", 11, align="center"),
    Column("واحد", "unit", 8, align="center"),
    Column("طبقه", "floor", 7, align="center"),
    Column("عنوان واحد", "unit_title", 20),
    Column("نیم‌روز", "half", 14, align="center"),
    Column("ساعت شروع", "start", 10, align="center"),
    Column("ساعت پایان", "end", 10, align="center"),
    Column("وضعیت سانس", "status", 16, align="center"),
    Column("توضیح", "detail", 34),
)

_UNIT_COLUMNS: tuple[Column, ...] = (
    Column("واحد", "unit", 8, align="center"),
    Column("طبقه", "floor", 7, align="center"),
    Column("عنوان واحد", "unit_title", 22),
    Column("نوع", "unit_type", 12, align="center"),
    Column("ظرفیت", "capacity", 9, align="center"),
    Column("نیم‌روز کل", "total", 11, align="center"),
    Column("اشغال", "occupied", 9, align="center"),
    Column("آزاد", "free", 9, align="center"),
    Column("بافر نظافت", "buffer", 12, align="center"),
    Column("غیرقابل‌رزرو", "blocked", 13, align="center"),
    Column("نرخ اشغال", "occupancy", 12, "0.0%", align="center"),
    Column("شب‌های رزروشده", "nights", 15, align="center"),
)

_DAILY_COLUMNS: tuple[Column, ...] = (
    Column("تاریخ", "day", 12, align="center"),
    Column("تاریخ شمسی", "date_jalali", 13, align="center"),
    Column("روز هفته", "weekday", 11, align="center"),
    Column("پایان هفته", "is_weekend", 12, align="center"),
    Column("سطح تقاضا", "demand", 13, align="center"),
    Column("نیم‌روز کل", "total", 11, align="center"),
    Column("اشغال", "occupied", 9, align="center"),
    Column("آزاد", "free", 9, align="center"),
    Column("نرخ اشغال", "occupancy", 12, "0.0%", align="center"),
    Column("ورودی", "arrivals", 9, align="center"),
    Column("خروجی", "departures", 9, align="center"),
)

_ACTION_COLUMNS: tuple[Column, ...] = (
    Column("زمان", "moment", 17, align="center"),
    Column("کارمند", "actor", 20),
    Column("کنش", "kind", 14, align="center"),
    Column("کد رزرو", "code", 13, align="center"),
    Column("واحد", "unit", 8, align="center"),
    Column("از وضعیت", "from_status", 14, align="center"),
    Column("به وضعیت", "to_status", 14, align="center"),
    Column("دلیل ثبت‌شده", "reason", 42),
)

_SMS_COLUMNS: tuple[Column, ...] = (
    Column("زمان", "moment", 17, align="center"),
    Column("قالب پیام", "template", 20, align="center"),
    Column("تلفن", "phone", 16, align="center"),
    Column("وضعیت", "status", 16, align="center"),
    Column("ارائه‌دهنده", "provider", 14, align="center"),
    Column("تلاش‌ها", "attempts", 9, align="center"),
)


def _jalali_stamp(moment: datetime) -> str:
    """مُهر زمانی شمسی برای برگهٔ راهنما."""
    from django.utils import timezone as django_timezone

    from apps.core.jalali import clock_label, format_jalali

    local = django_timezone.localtime(moment)
    return f"{format_jalali(local.date(), numeric=True)} ساعت {clock_label(local)}"


def sheet_titles(workbook: Workbook) -> list[str]:
    """فهرست نام برگه‌ها — برای تست ساختار و برای مستندسازی."""
    return list(workbook.sheetnames)


def range_label(report_range: ReportRange) -> str:
    """برچسب کوتاه بازه — کمکی برای نام فایل و عنوان."""
    return report_range.label
